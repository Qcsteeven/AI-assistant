# backend/document_loader.py
from typing import List
import docx
import openpyxl
import pdfplumber
from docx.document import Document as _Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from datasketch import MinHash, MinHashLSH


class DocumentLoader:
    def __init__(self):
        self.unique_chunks = set()
        self.lsh = MinHashLSH(threshold=0.85, num_perm=128)

    def _get_minhash(self, text: str, num_perm: int = 128) -> MinHash:
        m = MinHash(num_perm=num_perm)
        for word in text.split():
            m.update(word.encode('utf8'))
        return m

    def _filter_and_add_chunks(self, new_chunks: List[str]) -> List[str]:
        filtered_chunks = []
        for chunk in new_chunks:
            norm_chunk = chunk.strip()
            if not norm_chunk or norm_chunk in self.unique_chunks:
                continue
            m = self._get_minhash(norm_chunk)
            if not self.lsh.query(m):  # если нет похожего чанка
                self.lsh.insert(norm_chunk, m)
                self.unique_chunks.add(norm_chunk)
                filtered_chunks.append(norm_chunk)
        return filtered_chunks

    def _append_buffer(self, buffer: str, chunks: List[str]) -> str:
        if buffer.strip():
            chunks.append(buffer.strip())
        return ""

    def load_docx_chunks(self, path: str, chunk_size: int = 2000) -> List[str]:
        def iter_block_items(parent):
            if isinstance(parent, _Document):
                parent_elm = parent.element.body
            elif isinstance(parent, _Cell):
                parent_elm = parent._tc
            else:
                raise ValueError("Unsupported parent type")

            for child in parent_elm.iterchildren():
                if child.tag.endswith('}p'):
                    yield Paragraph(child, parent)
                elif child.tag.endswith('}tbl'):
                    yield docx.table.Table(child, parent)

        doc = docx.Document(path)
        chunks = []
        buffer = ""

        for block in iter_block_items(doc):
            if isinstance(block, docx.text.paragraph.Paragraph):
                text = block.text.strip()
                if text:
                    if len(buffer) + len(text) < chunk_size:
                        buffer += " " + text
                    else:
                        buffer = self._append_buffer(buffer, chunks)
                        buffer = text
            elif isinstance(block, docx.table.Table):
                for row in block.rows:
                    row_texts = []
                    for cell in row.cells:
                        cell_text = cell.text.strip().replace('\n', ' ')
                        if cell_text:
                            row_texts.append(cell_text)
                    row_line = " | ".join(row_texts)
                    if row_line:
                        if len(buffer) + len(row_line) < chunk_size:
                            buffer += " " + row_line
                        else:
                            buffer = self._append_buffer(buffer, chunks)
                            buffer = row_line

        self._append_buffer(buffer, chunks)
        return self._filter_and_add_chunks(chunks)

    def load_pdf_chunks(self, path: str, chunk_size: int = 1000) -> List[str]:
        with pdfplumber.open(path) as pdf:
            chunks = []
            buffer = ""

            for page in pdf.pages:
                text = page.extract_text() or ""
                for line in text.strip().split('\n'):
                    line = line.strip()
                    if line:
                        if len(buffer) + len(line) < chunk_size:
                            buffer += " " + line
                        else:
                            buffer = self._append_buffer(buffer, chunks)
                            buffer = line
            self._append_buffer(buffer, chunks)
            return self._filter_and_add_chunks(chunks)

    def load_xlsx_chunks(self, path: str, chunk_size: int = 1000) -> List[str]:
        wb = openpyxl.load_workbook(path)
        sheet = wb.active
        chunks = []
        buffer = ""

        for row in sheet.iter_rows():
            row_text = " | ".join(str(cell.value).strip() for cell in row if cell.value)
            if row_text:
                if len(buffer) + len(row_text) < chunk_size:
                    buffer += " " + row_text
                else:
                    buffer = self._append_buffer(buffer, chunks)
                    buffer = row_text

        self._append_buffer(buffer, chunks)
        return self._filter_and_add_chunks(chunks)