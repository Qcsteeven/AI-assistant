# backend/app.py
from typing import List, Optional, Tuple
import faiss
import numpy as np
import os
import pickle
import docx
import openpyxl
import pdfplumber
from openai import OpenAI
from docx.document import Document as _Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from datasketch import MinHash, MinHashLSH

class VectorBackend:
    _instance = None

    def __new__(cls):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
            cls._instance.__init__()
        return cls._instance

    def __init__(self):
        if not hasattr(self, 'is_initialized'):
            self.client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
            self.index: Optional[faiss.Index] = None
            self.embeddings: List[List[float]] = []
            self.chunks: List[str] = []
            self.unique_chunks = set()  # <--- Новое
            self.lsh = MinHashLSH(threshold=0.85, num_perm=128)  # <--- Новое
            self.is_initialized = True

    def _get_minhash(self, text: str, num_perm: int = 128) -> MinHash:
        m = MinHash(num_perm=num_perm)
        for word in text.split():
            m.update(word.encode('utf8'))
        return m

    def _filter_and_add_chunks(self, new_chunks: List[str]):
        for chunk in new_chunks:
            norm_chunk = chunk.strip()
            if not norm_chunk or norm_chunk in self.unique_chunks:
                continue
            m = self._get_minhash(norm_chunk)
            if not self.lsh.query(m):  # если нет похожего чанка
                self.lsh.insert(norm_chunk, m)
                self.unique_chunks.add(norm_chunk)
                self.chunks.append(norm_chunk)

    def _append_buffer(self, buffer: str, chunks: List[str]) -> str:
        if buffer.strip():
            chunks.append(buffer.strip())
        return ""

    def load_docx_chunks(self, path: str, chunk_size: int = 2000) -> List[str]:
        def iter_block_items(parent):
            if isinstance(parent, _Document):
                parent_elm = parent.element.body
            elif isinstance(parent, _Cell):
                parent_elm = parent._tc
            else:
                raise ValueError("Unsupported parent type")

            for child in parent_elm.iterchildren():
                if child.tag.endswith('}p'):
                    yield Paragraph(child, parent)
                elif child.tag.endswith('}tbl'):
                    yield docx.table.Table(child, parent)

        doc = docx.Document(path)
        chunks = []
        buffer = ""

        for block in iter_block_items(doc):
            if isinstance(block, docx.text.paragraph.Paragraph):
                text = block.text.strip()
                if text:
                    if len(buffer) + len(text) < chunk_size:
                        buffer += " " + text
                    else:
                        buffer = self._append_buffer(buffer, chunks)
                        buffer = text
            elif isinstance(block, docx.table.Table):
                for row in block.rows:
                    row_texts = []
                    for cell in row.cells:
                        # Собираем весь текст в ячейке, включая переносы строк
                        cell_text = cell.text.strip().replace('\n', ' ')
                        if cell_text:
                            row_texts.append(cell_text)
                    row_line = " | ".join(row_texts)
                    if row_line:
                        if len(buffer) + len(row_line) < chunk_size:
                            buffer += " " + row_line
                        else:
                            buffer = self._append_buffer(buffer, chunks)
                            buffer = row_line

        self._append_buffer(buffer, chunks)
        self._filter_and_add_chunks(chunks)  # <--- заменено
        return chunks

    def load_pdf_chunks(self, path: str, chunk_size: int = 1000) -> List[str]:
        with pdfplumber.open(path) as pdf:
            chunks = []
            buffer = ""

            for page in pdf.pages:
                text = page.extract_text() or ""
                for line in text.strip().split('\n'):
                    line = line.strip()
                    if line:
                        if len(buffer) + len(line) < chunk_size:
                            buffer += " " + line
                        else:
                            buffer = self._append_buffer(buffer, chunks)
                            buffer = line
            self._append_buffer(buffer, chunks)
            self._filter_and_add_chunks(chunks)  # <--- заменено
            return chunks

    def load_xlsx_chunks(self, path: str, chunk_size: int = 1000) -> List[str]:
        wb = openpyxl.load_workbook(path)
        sheet = wb.active
        chunks = []
        buffer = ""

        for row in sheet.iter_rows():
            row_text = " | ".join(str(cell.value).strip() for cell in row if cell.value)
            if row_text:
                if len(buffer) + len(row_text) < chunk_size:
                    buffer += " " + row_text
                else:
                    buffer = self._append_buffer(buffer, chunks)
                    buffer = row_text

        self._append_buffer(buffer, chunks)
        self._filter_and_add_chunks(chunks)  # <--- заменено
        return chunks

    def load_all_documents(self, folder_path: str):
        for file in os.listdir(folder_path):
            full_path = os.path.join(folder_path, file)
            if os.path.isfile(full_path):
                ext = os.path.splitext(file)[1].lower()
                print(f"📄 Обработка {file}...")
                if ext == '.docx':
                    self.load_docx_chunks(full_path)
                elif ext == '.pdf':
                    self.load_pdf_chunks(full_path)
                elif ext == '.xlsx':
                    self.load_xlsx_chunks(full_path)

    def get_embedding(self, text: str) -> List[float]:
        response = self.client.embeddings.create(
            input=[text],
            model="text-embedding-ada-002"
        )
        return response.data[0].embedding

    def build_faiss_index(self):
        self.embeddings = [self.get_embedding(chunk) for chunk in self.chunks]
        dim = len(self.embeddings[0])
        self.index = faiss.IndexFlatL2(dim)
        self.index.add(np.array(self.embeddings).astype("float32"))
        return self.index

    def save_index(self, index_file: str, embeddings_file: str, chunks_file: str):
        if self.index:
            faiss.write_index(self.index, index_file)
            with open(embeddings_file, 'wb') as ef:
                pickle.dump(self.embeddings, ef)
            with open(chunks_file, 'wb') as cf:
                pickle.dump(self.chunks, cf)

    def load_index(self, index_file: str, embeddings_file: str, chunks_file: str):
        self.index = faiss.read_index(index_file)
        with open(embeddings_file, 'rb') as ef:
            self.embeddings = pickle.load(ef)
        with open(chunks_file, 'rb') as cf:
            self.chunks = pickle.load(cf)

    def search_similar_chunks(self, query: str, k: int = 3) -> List[str]:
        if not self.index:
            raise ValueError("Индекс не инициализирован")
        query_vec = np.array([self.get_embedding(query)]).astype("float32")
        _, indices = self.index.search(query_vec, k)
        return [self.chunks[i] for i in indices[0]]

    def chat(self, memory_limit: int = 5):
        if not self.index:
            raise ValueError("Сначала вызовите build_faiss_index()")

        print("💬 Чат с ИИ. Напиши 'exit' для выхода.")
        history = []

        while True:
            query = input("\nТы: ")
            if query.lower() == "exit":
                break

            top_chunks = self.search_similar_chunks(query)
            context = "\n---\n".join(top_chunks)
            history_context = "\n".join(f"Вопрос: {q}\nОтвет: {a}" for q, a in history[-memory_limit:])

            prompt = (
                "Ты помощник, который отвечает по документу.\n"
                f"История:\n{history_context}\n\n"
                f"Контекст:\n{context}\n\n"
                f"Вопрос: {query}\nОтвет:"
            )

            response = self.client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": prompt}]
            )
            answer = response.choices[0].message.content.strip()
            print("\nGPT:", answer)
            history.append((query, answer))
